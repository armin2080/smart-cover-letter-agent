"""CV text extraction and heuristic section mapping.

Supports PDF (via ``pypdf``) and DOCX (via ``python-docx``). The mapping is
deliberately conservative: it only fills sections it can positively identify and
never invents content. Everything it produces is editable by the user on the
profile screen.
"""

from __future__ import annotations

import logging
import re
from pathlib import Path

logger = logging.getLogger(__name__)

# Heading aliases → section key. Matched case-insensitively against short lines.
SECTION_ALIASES: dict[str, tuple[str, ...]] = {
    "personal": ("contact", "personal details", "personal information", "details"),
    "summary": (
        "summary",
        "professional summary",
        "profile",
        "about",
        "about me",
        "objective",
        "career objective",
    ),
    "experience": (
        "experience",
        "work experience",
        "professional experience",
        "employment",
        "employment history",
        "work history",
    ),
    "education": ("education", "academic background", "academics", "studies"),
    "skills": (
        "skills",
        "technical skills",
        "core skills",
        "core competencies",
        "competencies",
        "expertise",
        "technologies",
    ),
    "projects": ("projects", "selected projects", "personal projects", "portfolio"),
    "languages": ("languages", "language skills"),
    "certificates": (
        "certificates",
        "certifications",
        "certification",
        "licenses",
        "awards",
        "achievements",
    ),
}

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(r"(?:\+?\d[\d\s().\-/]{7,}\d)")
NAME_RE = re.compile(r"^[A-Z][A-Za-z'\-]+(?:\s+[A-Z][A-Za-z'\-]+){1,3}$")


def extract_text(path: str | Path) -> str:
    """Return the plain text of a PDF or DOCX file. Empty string on failure."""

    path = Path(path)
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            return _extract_pdf(path)
        if suffix == ".docx":
            return _extract_docx(path)
    except Exception:  # pragma: no cover - defensive, depends on file contents
        logger.warning("Could not extract text from %s", path, exc_info=True)
    return ""


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    chunks: list[str] = []
    for page in reader.pages:
        try:
            chunks.append(page.extract_text() or "")
        except Exception:  # pragma: no cover
            continue
    return "\n".join(chunks)


def _extract_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    chunks = [paragraph.text for paragraph in document.paragraphs]
    # Include simple tables, which CVs frequently use for skills and dates.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))
    return "\n".join(chunks)


def _normalise(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _heading_key(line: str) -> str | None:
    """Return a section key if `line` looks like a section heading."""

    candidate = line.strip().strip(":").strip()
    if not candidate or len(candidate) > 60 or len(candidate.split()) > 6:
        return None
    # Headings rarely end with sentence punctuation.
    if candidate.endswith((".", ",", ";")):
        return None

    lowered = candidate.lower()
    for key, aliases in SECTION_ALIASES.items():
        for alias in aliases:
            if lowered == alias or lowered.startswith(f"{alias} "):
                return key
    return None


def parse_sections(text: str) -> dict[str, str]:
    """Split extracted CV text into sections keyed by profile section key."""

    text = _normalise(text)
    if not text:
        return {}

    buckets: dict[str, list[str]] = {}
    current: str | None = None
    preamble: list[str] = []

    for line in text.split("\n"):
        key = _heading_key(line)
        if key:
            current = key
            buckets.setdefault(current, [])
            continue

        if current is None:
            preamble.append(line)
        else:
            buckets[current].append(line)

    result: dict[str, str] = {}
    for key, lines in buckets.items():
        content = "\n".join(line for line in lines).strip()
        content = re.sub(r"\n{2,}", "\n", content).strip()
        if content:
            result[key] = content

    # Text before the first heading usually holds contact details and a summary.
    pre = "\n".join(preamble).strip()
    if pre and "personal" not in result:
        result["personal"] = re.sub(r"\n{2,}", "\n", pre).strip()

    return result


def parse_contact(text: str) -> dict[str, str]:
    """Best-effort extraction of contact fields from raw CV text."""

    text = _normalise(text)
    contact: dict[str, str] = {}

    email_match = EMAIL_RE.search(text)
    if email_match:
        contact["email"] = email_match.group(0)

    # Ignore phone-like matches that are part of a longer numeric token.
    phone_match = PHONE_RE.search(text)
    if phone_match:
        candidate = phone_match.group(0).strip()
        if 8 <= len(re.sub(r"\D", "", candidate)) <= 15:
            contact["phone"] = candidate

    for line in text.split("\n")[:8]:
        stripped = line.strip()
        if NAME_RE.match(stripped) and "email" not in stripped.lower():
            contact["full_name"] = stripped
            break

    return contact


def parse_cv(path: str | Path) -> dict:
    """Extract everything we can from a CV file.

    Returns a dict with `text`, `sections` and `contact`.
    """

    text = extract_text(path)
    return {
        "text": text,
        "sections": parse_sections(text),
        "contact": parse_contact(text),
    }
